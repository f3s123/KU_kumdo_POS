from flask import Flask, render_template, request, jsonify, redirect, url_for, send_file
import sqlite3
import json
from datetime import datetime
import openpyxl
from io import BytesIO
from flask_socketio import SocketIO

app = Flask(__name__)
app.config['SECRET_KEY'] = 'it3rNullsec!'
socketio = SocketIO(app)

DB_PATH = "kendo_bar.db"

def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA encoding = 'UTF-8'")

    def dict_factory(cursor, row):
        return {col[0]: row[idx] for idx, col in enumerate(cursor.description)}

    conn.row_factory = dict_factory
    conn.text_factory = str
    return conn

# 메인 페이지
@app.route('/')
def main():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT table_num, total_price, entrance_time, end_time FROM table_orders")
    rows = cursor.fetchall()

    tables = []
    for row in rows:
        table_num = row['table_num']
        total_price = row['total_price']
        entrance_time = row['entrance_time']

        if entrance_time and entrance_time != "0":
            try:
                entrance_dt = datetime.strptime(entrance_time, '%Y-%m-%d %H:%M:%S')
                elapsed = datetime.now() - entrance_dt
                elapsed_time_str = f"{elapsed.seconds // 3600:02}시 {(elapsed.seconds % 3600) // 60:02}분 {elapsed.seconds % 60:02}초"
            except ValueError:
                elapsed_time_str = "00시 00분 00초"
        else:
            elapsed_time_str = "00시 00분 00초"

        tables.append({
            'table_num': table_num,
            'total_price': total_price,
            'elapsed_time': elapsed_time_str
        })

    conn.close()
    return render_template('./main.html', tables=tables)

# 테이블 페이지
@app.route('/table/<int:table_num>')
def table(table_num):
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
    SELECT tbl_orders, total_price, entrance_time, memo
    FROM table_orders WHERE table_num = ?
    """, (table_num,))
    row = cursor.fetchone()
    if not row:
        return "Table not found", 404

    tbl_orders = row['tbl_orders']
    total_price = row['total_price']
    entrance_time = row['entrance_time']
    memo = row['memo']

    elapsed_seconds = 0
    if entrance_time and entrance_time != "0":
        try:
            entry_time = datetime.strptime(entrance_time, '%Y-%m-%d %H:%M:%S')
            now = datetime.now()
            elapsed = now - entry_time
            elapsed_seconds = int(elapsed.total_seconds())
            elapsed_time_str = f"{elapsed_seconds // 3600:02}시간 {(elapsed_seconds % 3600) // 60:02}분 {elapsed_seconds % 60:02}초"
        except ValueError:
            elapsed_time_str = "00시간 00분 00초"
    else:
        elapsed_time_str = "00시간 00분 00초"

    cursor.execute("SELECT menu_json FROM menu")
    menu_data = json.loads(cursor.fetchone()['menu_json'])

    conn.close()
    return render_template('./table.html', 
                           table_num=table_num,
                           total_price=total_price,
                           elapsed_time=elapsed_time_str,
                           elapsed_seconds=elapsed_seconds,
                           entrance_time=entrance_time if entrance_time and entrance_time != "0" else "",
                           menus=menu_data['menus'],
                           orders=json.loads(tbl_orders),
                           memo=memo)

# 주문 추가
@app.route('/order', methods=['POST'])
def add_order():
    try:
        data = request.json
        table_num = data['table_num']
        menu_name = data['menu_name']
        price = data['price']
        timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("""
        INSERT INTO orders (table_num, menu_name, price, time)
        VALUES (?, ?, ?, ?)
        """, (table_num, menu_name, price, timestamp))

        cursor.execute("SELECT tbl_orders, total_price FROM table_orders WHERE table_num = ?", (table_num,))
        row = cursor.fetchone()
        if row:
            tbl_orders = json.loads(row['tbl_orders'])
            total_price = row['total_price']
            tbl_orders[menu_name] = tbl_orders.get(menu_name, 0) + 1
            total_price += price

            cursor.execute("""
            UPDATE table_orders 
            SET tbl_orders = ?, total_price = ?
            WHERE table_num = ?
            """, (json.dumps(tbl_orders), total_price, table_num))

        conn.commit()
        conn.close()
        return jsonify({'message': 'Order added successfully'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# 주문 취소
@app.route('/cancel', methods=['POST'])
def cancel_order():
    try:
        data = request.json
        #order_id = data.get('order_id')  # 안전하게 가져오기
        table_num = data.get('table_num')
        menu_name = data.get('menu_name')
        price = data.get('price')
        timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        if not table_num or not menu_name or not price:
            return jsonify({'error': 'Invalid input data'}), 400

        conn = get_db_connection()
        cursor = conn.cursor()

        # 취소된 주문을 deleted_orders 테이블에 추가
        cursor.execute("""
        INSERT INTO deleted_orders (table_num, menu, order_time, delete_time)
        SELECT table_num, menu_name, time, ?
        FROM orders
        WHERE table_num = ? AND menu_name = ?
        """, (timestamp, table_num, menu_name))

        # table_orders에서 주문 정보 업데이트
        cursor.execute("SELECT tbl_orders, total_price FROM table_orders WHERE table_num = ?", (table_num,))
        row = cursor.fetchone()
        if row:
            tbl_orders = json.loads(row['tbl_orders'])
            total_price = row['total_price']

            if menu_name in tbl_orders and tbl_orders[menu_name] > 0:
                tbl_orders[menu_name] -= 1
                total_price -= price
                cursor.execute("""
                UPDATE table_orders 
                SET tbl_orders = ?, total_price = ?
                WHERE table_num = ?
                """, (json.dumps(tbl_orders), total_price, table_num))

        conn.commit()
        conn.close()
        return jsonify({'message': 'Order canceled successfully'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# 테이크아웃 페이지
@app.route('/takeout')
def takeout():
    conn = get_db_connection()
    c = conn.cursor()

    # menu_to 테이블에서 데이터 가져오기
    c.execute('SELECT menu_json FROM menu_to')
    result = c.fetchone()
    if not result:
        conn.close()
        return "메뉴 데이터가 없습니다.", 404

    menu_data = json.loads(result['menu_json'])

    # payments 테이블에서 가장 큰 테이블 번호 가져오기
    c.execute("SELECT MAX(table_num) as max_num FROM payments WHERE table_num >= 20")
    result = c.fetchone()
    max_table_num = result['max_num'] if result and result['max_num'] is not None else None

    # 테이크아웃 번호 설정
    if max_table_num is None or max_table_num < 20:
        takeout_number = 20  # 기본값 설정
    else:
        takeout_number = max_table_num + 1

    conn.close()
    return render_template('takeout.html', menus=menu_data['menus'], takeout_number=takeout_number)

@app.route('/submit-takeout-orders', methods=['POST'])
def submit_takeout_orders():
    try:
        data = request.json  # 클라이언트에서 전송된 주문 데이터
        orders = data.get('orders', {})  # 주문 데이터
        memo = data.get('memo', '')  # 메모 데이터
        conn = get_db_connection()
        c = conn.cursor()

        # 다음 테이크아웃 번호 가져오기
        c.execute("SELECT MAX(table_num) as max_num FROM payments WHERE table_num >= 20")
        result = c.fetchone()
        max_table_num = result['max_num'] if result and result['max_num'] is not None else None

        # 테이크아웃 번호 설정
        if max_table_num is None or max_table_num < 20:
            takeout_number = 20  # 기본값 설정
        else:
            takeout_number = max_table_num + 1
        print(f"takeout_number: {takeout_number}")

        # 주문 데이터 저장
        total_price = 0
        order_details = {}
        for menu_name, data in orders.items():
            count = data['count']
            price = data['price']
            total_price += count * price
            order_details[menu_name] = count
            for _ in range(count):
                c.execute('''
                    INSERT INTO orders (table_num, menu_name, price, time, etc)
                    VALUES (?, ?, ?, datetime('now', 'localtime'), 0)
                ''', (takeout_number, menu_name, price))

        # payments 테이블에 정보 저장
        payment_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        c.execute('''
            INSERT INTO payments (table_num, total_price, payment_time, memo, detail, entrance_time, end_time, used_time, used_seconds)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (takeout_number, total_price, payment_time, memo, json.dumps(order_details, ensure_ascii=False), '0', '0', '0', 0))

        conn.commit()
        conn.close()
        return jsonify({'message': '주문이 완료되었습니다.'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# 계산 완료
@app.route('/complete/<int:table_num>', methods=['POST'])
def complete_order(table_num):
    try:
        now = datetime.now()
        timestamp = now.strftime('%Y-%m-%d %H:%M:%S')

        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT tbl_orders, total_price, entrance_time, memo FROM table_orders WHERE table_num = ?", (table_num,))
        row = cursor.fetchone()

        if row:
            tbl_orders = row['tbl_orders']
            total_price = row['total_price']
            entrance_time = row['entrance_time']
            memo = row['memo'] if row['memo'] else "" # 메모 가져오기

            if entrance_time and entrance_time != "0":
                entrance_dt = datetime.strptime(entrance_time, '%Y-%m-%d %H:%M:%S')
                used_time = now - entrance_dt
                used_seconds = int(used_time.total_seconds())
                used_time_str = f"{used_seconds // 3600:02}시 {(used_seconds % 3600) // 60:02}분 {used_seconds % 60:02}초"

                cursor.execute("""
                INSERT INTO payments 
                (table_num, total_price, payment_time, memo, detail, entrance_time, end_time, used_time, used_seconds)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (table_num, total_price, timestamp, memo, tbl_orders, entrance_time, timestamp, used_time_str, used_seconds))

            reset_orders = {
                "타코야끼 (데리야끼)": 0,
                "타코야끼 (불닭)": 0,
                "야끼소바 (간장)": 0,
                "야끼소바 (불닭)": 0,
                "우삼겹숙주볶음": 0,
                "나가사키해물우동": 0,
                "흑당인절미 당고": 0,
                "황도": 0,
                "교자": 0,
                "메론소다": 0,
                "청포도 에이드": 0,
                "망고 에이드": 0,
                "아망추": 0,
                "선라이즈": 0,
                "로이 로저스": 0,
                "신데렐라": 0,
                "하이볼 키트": 0,
                "입장료 + 자릿세": 0,
                "콜키지 1L 미만": 0,
                "콜키지 1L 이상": 0
            }
            cursor.execute("""
            UPDATE table_orders 
            SET tbl_orders = ?, total_price = 0, entrance_time = '', end_time = '', memo = ''
            WHERE table_num = ?
            """, (json.dumps(reset_orders, ensure_ascii=False), table_num))

        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': 'Payment completed successfully'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# 시간 시작
@app.route('/start-timer/<int:table_num>', methods=['POST'])
def start_timer(table_num):
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT entrance_time FROM table_orders WHERE table_num = ?", (table_num,))
    row = cursor.fetchone()
    if row and (not row['entrance_time'] or row['entrance_time'] == "0"):
        cursor.execute("""
        UPDATE table_orders 
        SET entrance_time = ?
        WHERE table_num = ?
        """, (timestamp, table_num))
        conn.commit()

    conn.close()
    return jsonify({'success': True, 'timestamp': timestamp})

@app.route('/save-memo/<int:table_num>', methods=['POST'])
def save_memo(table_num):
    try:
        data = request.json
        memo = data.get('memo', '')  # 클라이언트에서 전달된 메모

        conn = get_db_connection()
        cursor = conn.cursor()

        # table_orders 테이블에 메모 저장
        cursor.execute("""
        UPDATE table_orders
        SET memo = ?
        WHERE table_num = ?
        """, (memo, table_num))

        conn.commit()
        conn.close()

        return jsonify({'message': 'Memo saved successfully'}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/order-menu-list')
def order_menu_list():
    # 고정된 대표 메뉴 목록 (화면에 표시할 메뉴들)
    menus = [
        {'id': 1, 'name': '타코야끼'},
        {'id': 2, 'name': '야끼소바'},
        {'id': 3, 'name': '우삼겹숙주볶음'},
        {'id': 4, 'name': '나가사키해물우동'},
        {'id': 5, 'name': '사이드'},
        {'id': 6, 'name': '음료'},
    ]
    return render_template('order_menu_list.html', menus=menus)

@app.route('/menu-orders/<category>')
def view_menu_orders(category):
    # 카테고리별 메뉴 이름 매핑
    menu_category_map = {
        "타코야끼": ["타코야끼 (데리야끼)", "타코야끼 (불닭)"],
        "야끼소바": ["야끼소바 (간장)", "야끼소바 (불닭)"],
        "우삼겹숙주볶음": ["우삼겹숙주볶음"],
        "나가사키해물우동": ["나가사키해물우동"],
        "사이드": ["흑당인절미 당고", "황도", "교자"],
        "음료": [
            "메론소다", "청포도 에이드", "망고 에이드", "아망추",
            "선라이즈", "로이 로저스", "신데렐라",
            "하이볼 키트"
        ]
    }

    if category not in menu_category_map:
        return f"존재하지 않는 카테고리입니다: {category}", 404

    menu_names = menu_category_map[category]
    conn = get_db_connection()
    cursor = conn.cursor()

    # 메뉴 데이터 가져오기
    cursor.execute("SELECT menu_json FROM menu")
    row = cursor.fetchone()
    if not row:
        conn.close()
        return "메뉴 데이터가 없습니다.", 404

    menu_data = json.loads(row['menu_json'])
    menus = [menu.copy() for menu in menu_data['menus'] if menu['name'] in menu_names]

    # 해당 메뉴들에 대한 주문 가져오기
    placeholders = ','.join(['?'] * len(menu_names))
    orders = cursor.execute(
        f"SELECT order_id, table_num, menu_name FROM orders WHERE menu_name IN ({placeholders})",
        menu_names
    ).fetchall()

    # 메뉴별로 여러 개의 주문 저장
    for menu in menus:
        menu['orders'] = [
            {'order_id': order['order_id'], 'table_num': order['table_num']}
            for order in orders if order['menu_name'] == menu['name']
        ]

    conn.close()
    return render_template('menu_orders.html', category=category, menus=menus)

@app.route('/complete-order', methods=['POST'])
def complete_order_one():
    order_id = request.form['order_id']
    menu_name = request.form['menu_name']
    table_num = request.form['table_num']
    category = request.form.get('category')  # POST 요청에서 category 값 가져오기
    done_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')  # 완료 시간 기록

    conn = get_db_connection()
    cursor = conn.cursor()

    # 완료된 주문을 done_orders 테이블에 추가
    cursor.execute("""
    INSERT INTO done_orders (table_num, menu, order_time, done_time)
    SELECT table_num, menu_name, time, ?
    FROM orders
    WHERE order_id = ?
    """, (done_time, order_id))

    # orders 테이블에서 해당 주문 삭제
    cursor.execute("""
    DELETE FROM orders
    WHERE order_id = ?
    """, (order_id,))

    conn.commit()
    conn.close()

    # category 값이 없으면 기본값 설정
    if not category:
        category = "기본 카테고리"  # 필요에 따라 기본값 설정

    return redirect(url_for('view_menu_orders', category=category))

@app.route('/done_orders')
def view_done_orders():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT table_num, menu, order_time, done_time FROM done_orders ORDER BY done_time DESC")
    done_orders = cur.fetchall()
    
    conn.commit()
    conn.close()

    return render_template('done_orders.html', done_orders=done_orders)

@app.route('/payments')
def view_payments():
    db = get_db_connection()
    cur = db.cursor()
    
    # payments 테이블에서 데이터 가져오기
    cur.execute("SELECT * FROM payments")
    rows = cur.fetchall()

    payments = []
    total_revenue = 0  # 총 결제 금액
    for row in rows:
        payments.append({
            'table_num': row['table_num'],
            'total_price': row['total_price'],
            'payment_time': row['payment_time'],
            'memo': row['memo'],
            'detail': json.loads(row['detail']) if row['detail'] else {},
            'used_time': row['used_time']
        })
        total_revenue += row['total_price']

    db.commit()
    db.close()

    return render_template('payments.html', payments=payments, total_revenue=total_revenue)

@app.route('/export-payments')
def export_payments():
    conn = get_db_connection()
    cursor = conn.cursor()

    # payments 테이블에서 데이터 가져오기
    cursor.execute("SELECT * FROM payments")
    rows = cursor.fetchall()

    # 메뉴 가격 정보를 가져오기
    cursor.execute("SELECT menu_json FROM menu")
    menu_data = json.loads(cursor.fetchone()['menu_json'])
    menu_prices = {menu['name']: menu['price'] for menu in menu_data['menus']}

    # 엑셀 파일 생성
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments"

    # 헤더 추가
    headers = ["테이블 번호", "총 결제 금액", "입장 시간", "결제 시각", "사용 시간", "메모", "상세 주문"]
    ws.append(headers)

    # 데이터 추가
    for row in rows:
        detail = json.loads(row['detail'])
        detail_str = ", ".join([
            f"{menu_name}({count}개 * {menu_prices.get(menu_name, 0)} = ₩{count * menu_prices.get(menu_name, 0)})"
            for menu_name, count in detail.items() if count > 0
        ])
        ws.append([
            row['table_num'],
            row['total_price'],
            row['entrance_time'], 
            row['payment_time'],
            row['used_time'],
            row['memo'],
            detail_str
        ])

    # 엑셀 파일을 메모리에 저장
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    conn.close()

    # 엑셀 파일 다운로드
    return send_file(output, as_attachment=True, download_name="payments.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=502, debug=True)
